from __future__ import annotations

import textwrap
import uuid
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas

from .resume_parser import parse_resume_text

TEMPLATE_STYLES = {
    "classic": {
        "font": "Calibri",
        "name_size": 20,
        "body_size": 11,
        "heading_color": RGBColor(0x22, 0x22, 0x22),
        "pdf_font": "Helvetica",
    },
    "modern": {
        "font": "Georgia",
        "name_size": 22,
        "body_size": 11,
        "heading_color": RGBColor(0x00, 0x5A, 0x9C),
        "pdf_font": "Times-Roman",
    },
    "minimal": {
        "font": "Arial",
        "name_size": 19,
        "body_size": 10,
        "heading_color": RGBColor(0x33, 0x33, 0x33),
        "pdf_font": "Courier",
    },
}


def export_resume_file(
    resume_text: str,
    style: str,
    output_format: str,
    output_dir: Path,
    base_filename: str = "fixed_resume",
) -> Path:
    style = style.lower()
    output_format = output_format.lower()
    if style not in TEMPLATE_STYLES:
        raise ValueError(f"Unknown style '{style}'. Choose from: {', '.join(TEMPLATE_STYLES)}")

    output_dir.mkdir(parents=True, exist_ok=True)
    safe_name = f"{base_filename}_{style}_{uuid.uuid4().hex[:8]}"
    output_path = output_dir / f"{safe_name}.{output_format}"

    data = parse_resume_text(resume_text)
    if output_format == "docx":
        _export_docx(data, style, output_path)
    elif output_format == "pdf":
        _export_pdf(data, style, output_path)
    else:
        raise ValueError("output_format must be 'docx' or 'pdf'")

    return output_path


def _export_docx(data, style: str, output_path: Path) -> None:
    style_cfg = TEMPLATE_STYLES[style]
    doc = Document()

    for section in doc.sections:
        section.top_margin = Pt(40)
        section.bottom_margin = Pt(40)
        section.left_margin = Pt(40)
        section.right_margin = Pt(40)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = style_cfg["font"]
    normal_style.font.size = Pt(style_cfg["body_size"])

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(data.name)
    title_run.bold = True
    title_run.font.size = Pt(style_cfg["name_size"])

    contact_bits = [bit for bit in [data.email, data.phone] if bit]
    contact_bits.extend(data.links[:2])
    if contact_bits:
        contact = doc.add_paragraph(" | ".join(contact_bits))
        contact.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for section_name, lines in data.sections.items():
        heading = doc.add_paragraph()
        heading_run = heading.add_run(section_name.upper())
        heading_run.bold = True
        heading_run.font.color.rgb = style_cfg["heading_color"]
        heading_run.font.size = Pt(style_cfg["body_size"] + 1)

        for line in lines:
            text = line.lstrip("-*\u2022 ").strip()
            if not text:
                continue
            p = doc.add_paragraph(text, style="List Bullet")
            p.paragraph_format.space_after = Pt(4)

    doc.save(str(output_path))


def _export_pdf(data, style: str, output_path: Path) -> None:
    style_cfg = TEMPLATE_STYLES[style]
    c = canvas.Canvas(str(output_path), pagesize=LETTER)
    width, height = LETTER
    x = 52
    y = height - 50

    c.setFont(style_cfg["pdf_font"], 18)
    c.drawString(x, y, data.name)
    y -= 20

    contact_bits = [bit for bit in [data.email, data.phone] if bit]
    contact_bits.extend(data.links[:2])
    if contact_bits:
        c.setFont(style_cfg["pdf_font"], 10)
        c.drawString(x, y, " | ".join(contact_bits))
        y -= 18

    for section_name, lines in data.sections.items():
        if y < 90:
            c.showPage()
            y = height - 50
        c.setFont(style_cfg["pdf_font"], 12)
        c.drawString(x, y, section_name.upper())
        y -= 14
        c.setFont(style_cfg["pdf_font"], 10)
        for line in lines:
            clean = line.lstrip("-*\u2022 ").strip()
            if not clean:
                continue
            wrapped = textwrap.wrap(clean, width=95)
            for chunk in wrapped:
                if y < 70:
                    c.showPage()
                    y = height - 50
                    c.setFont(style_cfg["pdf_font"], 10)
                c.drawString(x + 8, y, f"- {chunk}")
                y -= 12
        y -= 6

    c.save()

